"""Generate the speech, presenter notes, source report, static fallback and PDF.

Inputs: exported content.js data and actual audience captures, supplied with --work.
Dependencies for document/PDF generation: python-docx, Pillow, reportlab.
"""
from pathlib import Path
import argparse, json, re, html, shutil, math
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

parser = argparse.ArgumentParser()
parser.add_argument('--work', required=True)
args = parser.parse_args()
work = Path(args.work)
project = Path(__file__).resolve().parents[1]
out = work / 'deliverables'
out.mkdir(exist_ok=True)
D = json.loads((work/'qa/keynote-data.json').read_text())
captures = json.loads((work/'qa/frames/screens.json').read_text())
SITE = 'https://antonov-ulyanovsk-keynote-2026.alyshagf.chatgpt.site'
REPO = 'https://github.com/ANTOVIUM/sgua/tree/ulyanovsk-keynote-20261007/ulyanovsk-keynote-2026'

def spoken(s):
    for a,b in [('ГК РФ','Гражданского кодекса Российской Федерации'),('434.1','четыреста тридцать четыре точка один'),('179','сто семьдесят девять'),('2016','две тысячи шестнадцатого'),('№','номер'),('de lege ferenda','де леге ференда')]:
        s=s.replace(a,b)
    return re.sub(r'([А-Я])\.([А-Я])\.',r'\1 \2',s)

steps=[]
for si,s in enumerate(D['scenes']):
    for ti,t in enumerate(s['steps']):
        steps.append(dict(t,scene=si,step=ti,words=len(spoken(t['speech']).split())))
pauses=sum(t['pause'] for t in steps)
word_total=sum(t['words'] for t in steps)
written=sum(len(t['speech'].split()) for t in steps)
cursor=0
for t in steps:
    t['start']=cursor
    t['duration']=(600-pauses)*t['words']/word_total+t['pause']
    cursor+=t['duration']
def tc(s):
    n=max(0,math.floor(s+1e-6))
    return f'{n//60:02}:{n%60:02}'

def field(p,name):
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),name);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p._p.append(f)

def docbase(title, body_size=14, line=1.5):
    d=Document(); sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
    sec.left_margin=Cm(2.3);sec.right_margin=Cm(2)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
        st._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Times New Roman')
        fonts=st._element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():del fonts.attrib[attribute]
        pp=st._element.get_or_add_pPr()
        for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(body_size)
    normal.paragraph_format.line_spacing=line
    normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    normal.paragraph_format.first_line_indent=Cm(1)
    d.styles['Title'].font.size=Pt(20);d.styles['Title'].font.bold=True
    d.styles['Title'].paragraph_format.space_after=Pt(14)
    d.styles['Title'].paragraph_format.keep_with_next=True
    d.styles['Heading 1'].font.size=Pt(16);d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(16)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    d.styles['Heading 2'].font.size=Pt(13);d.styles['Heading 2'].font.bold=True
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.paragraph_format.first_line_indent=Cm(0)
        st.paragraph_format.keep_with_next=True
    d.add_paragraph(title,'Title')
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field(f,'PAGE');f.paragraph_format.first_line_indent=Cm(0)
    d.core_properties.author='Антонов Михаил Александрович'
    d.core_properties.title=title
    return d

def p(d,s,bold=False,small=False):
    q=d.add_paragraph();q.paragraph_format.first_line_indent=Cm(0 if bold or small else 1)
    r=q.add_run(s);r.bold=bold
    if bold and not small:q.paragraph_format.keep_with_next=True
    if small:r.font.size=Pt(11)
    return q

def link(d,label,url):
    q=d.add_paragraph();q.paragraph_format.first_line_indent=Cm(0)
    q.add_run(label+'\n').bold=True
    relation=d.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),relation)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');rp.append(font)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');rp.append(size)
    r.append(rp);t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);q._p.append(h)
    return q

# The speech body comes only from the same step data shown in presenter mode.
speech=docbase('Место презумпции недобросовестности в механизме защиты гражданских прав')
p(speech,'Антонов Михаил Александрович\nПреподаватель кафедры гражданского права\nФГБОУ ВО «Саратовская государственная юридическая академия»',small=True)
p(speech,'Поволжский юридический конгресс\nУльяновск 10 октября 2026 года',small=True)
p(speech,f'Текст устного выступления. {written} слов. Плановый регламент 10 минут с {pauses} секундами пауз. Служебные строки и заголовок не произносятся.',small=True)
for t in steps:p(speech,t['speech'])
speech.save(out/'Antonov_Speech_10min.docx')

notes=docbase('Заметки докладчика для интерактивного выступления',12.5,1.15)
p(notes,'Антонов Михаил Александрович\nУльяновск 10 октября 2026 года',small=True)
p(notes,'Восемь сцен и 26 смысловых шагов образуют основную десятиминутную линию. Резерв из десяти сцен открывается отдельно и не входит в регламент. Текст ниже совпадает с экраном докладчика.',small=True)
p(notes,f'Расчетный темп: {word_total/(568/60):.1f} слова в минуту после раскрытия сокращений; {pauses} секунды выделены на паузы. Таймкоды обозначают плановые окна, а не измеренное выступление автора.',small=True)
notes.add_heading('Запуск и управление',1)
p(notes,'Для выступления без сети открой standalone.html. Клавиша P открывает отдельное окно докладчика; перенеси его на экран ноутбука, а аудиторию - на проектор. Два окна следует открывать через P или через кнопку «Открыть аудиторию», особенно при запуске с диска.',small=True)
p(notes,'Если браузер ограничивает взаимодействие локальных файлов, запусти python start-local.py в папке dist. Локальный сервер работает на 127.0.0.1 и не требует интернета. Не закрывай терминал во время выступления.',small=True)
p(notes,'До первой фразы: Home, сброс таймера в панели докладчика, F, затем T. F11 - резервный полный экран, если браузер запрещает Fullscreen API. Проверяй фокус окна перед нажатием кликера.',small=True)
controls=[('ArrowRight Space Enter PageDown','Следующий смысловой шаг'),('ArrowLeft PageUp','Предыдущий смысловой шаг'),('F','Полный экран'),('P','Окно докладчика'),('T','Старт или пауза таймера'),('A или Q затем 1-9 или 0','Резервная тема 1-10'),('Esc','Закрыть меню или вернуться из резерва'),('M','Сокращенное движение'),('G','Плоская геометрия'),('B или точка','Черный экран'),('Home','Первая сцена')]
for a,b in controls:p(notes,a+' - '+b,small=True)
notes.add_page_break()
notes.add_heading('Основная линия',1)
for si,s in enumerate(D['scenes']):
    ts=[t for t in steps if t['scene']==si]
    notes.add_heading(f'{si+1:02} {s["title"]}',1)
    p(notes,f'План {tc(ts[0]["start"])}-{tc(ts[-1]["start"]+ts[-1]["duration"])}. Опора в статье: {s["article"]}.',small=True)
    for t in ts:
        notes.add_heading(f'Шаг {t["step"]+1} {t["label"]}',2)
        q=p(notes,f'Окно {tc(t["start"])}-{tc(t["start"]+t["duration"])}. Пауза {t["pause"]} сек.',small=True)
        q.paragraph_format.keep_with_next=True
        q=p(notes,t['speech']);q.paragraph_format.keep_with_next=True
        q.paragraph_format.keep_together=True
        p(notes,('Остаться на финальном экране. ' if t is steps[-1] else 'После фразы '+t['cue']+' нажать один раз. ')+f'Пауза {t["pause"]} сек.',bold=True,small=True)
notes.add_heading('Резерв для вопросов',1)
for i,q in enumerate(D['qa']):
    notes.add_heading(f'{i+1:02} {q["title"]}',2)
    trigger=p(notes,f'A затем {i+1 if i<9 else 0}. '+q['ref'],small=True);trigger.paragraph_format.keep_with_next=True
    para=p(notes,q['body']);para.paragraph_format.keep_together=True
p(notes,'При входе в резерв таймер приостанавливается. Esc возвращает к сохраненному шагу основной линии; продолжение таймера запускается отдельно клавишей T.',small=True)
notes.save(out/'Antonov_Presenter_Notes.docx')

# Pure HTML/SVG scaling works without JavaScript and retains selectable text.
static=project/'dist/static';static.mkdir(exist_ok=True)
for i,s in enumerate(captures):
    name='index.html' if i==0 else (f'step-{i+1:02}.html' if i<26 else f'qa-{i-25:02}.html')
    names=['index.html']+[f'step-{n:02}.html' for n in range(2,27)]+[f'qa-{n:02}.html' for n in range(1,11)]
    prev=names[max(0,i-1)];nxt=names[min(35,i+1)]
    cls=' '.join(x for x in s['cls'].split() if x not in ['motion','scene-enter'])
    bg='#f1eee7' if 'paper' in cls else '#151718'
    title=D['scenes'][int(s['scene'])]['title'] if i<26 else D['qa'][i-26]['title']
    markup=s['html']
    page=f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)} · статическая версия</title><link rel="stylesheet" href="../style.css"><style>html,body{{background:{bg}}}main{{width:100vw;height:100vh}}main>svg{{display:block;width:100%;height:100%}}.stage{{position:relative;transform:none!important}}.static-nav{{position:fixed;bottom:5px;left:50%;transform:translateX(-50%);z-index:5;opacity:.42;display:flex;gap:18px;background:{bg};color:{"#151718" if "paper" in cls else "#f1eee7"};padding:4px 12px;font-size:12px}}.static-nav:hover,.static-nav:focus-within{{opacity:1}}</style></head><body class="reduced"><main><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" preserveAspectRatio="xMidYMid meet"><foreignObject width="1280" height="720"><div xmlns="http://www.w3.org/1999/xhtml" class="{cls}" aria-label="{html.escape(title)}">{markup}</div></foreignObject></svg></main><nav class="static-nav" aria-label="Статическое управление"><a href="{prev}" accesskey="p">Назад</a><a href="{nxt}" accesskey="n">Далее</a><a href="index.html">Начало</a><a href="qa-01.html">Резерв</a><a href="../fallback.pdf">PDF</a></nav></body></html>'''
    (static/name).write_text(page)

# Rendered audience composition, one semantic state per PDF page.
pdf_path=out/'Antonov_Keynote_Fallback.pdf'
c=canvas.Canvas(str(pdf_path),pagesize=(960,540),pageCompression=1)
c.setTitle('Место презумпции недобросовестности в механизме защиты гражданских прав')
c.setAuthor('Антонов Михаил Александрович')
c.setSubject('26 основных состояний и 10 резервных сцен')
cropped=work/'qa/cropped';cropped.mkdir(exist_ok=True)
for i,s in enumerate(captures):
    im=Image.open(work/'qa/frames'/s['file']).convert('RGB')
    r=s['rect']; x,y,w,h=r['x'],r['y'],r['width'],r['height']
    im=im.crop((round(x),round(y),round(x+w),round(y+h)))
    crop_path=cropped/s['file'];im.save(crop_path,quality=95)
    factor=960/w
    full=Image.open(work/'qa/frames'/s['file'])
    c.drawImage(str(work/'qa/frames'/s['file']),-x*factor,(y+h-full.height)*factor,width=full.width*factor,height=full.height*factor)
    bookmark=f'frame-{i+1:02}';c.bookmarkPage(bookmark)
    if i<26:
        title=f'{int(s["scene"])+1:02} {D["scenes"][int(s["scene"])]["title"]} шаг {int(s["step"])+1}'
    else:title=f'Резерв {i-25:02} {D["qa"][i-26]["title"]}'
    c.addOutlineEntry(title,bookmark,level=0)
    c.showPage()
c.save();shutil.copyfile(pdf_path,project/'dist/fallback.pdf')

# Test report is updated after publication with returned URLs and deployment ID.
report=docbase('Проверка интерактивного научного выступления',12,1.15)
p(report,'Антонов Михаил Александрович\nПроверка выполнена 7 октября 2026 года',small=True)
p(report,'Комплект содержит интерактивный экран аудитории, отдельный экран докладчика, восемь основных сцен с 26 смысловыми состояниями, десять резервных тем, автономный HTML, статическую версию без JavaScript и PDF из 36 страниц. Речь получена непосредственно из статьи; научная конструкция и авторское предложение сохранены.')
report.add_heading('Авторский текст и хронометраж',1)
p(report,f'Объем произносимого текста - {written} слов. После раскрытия сокращений в расчетной модели - {word_total} слов. При темпе {word_total/((600-pauses)/60):.1f} слова в минуту речь занимает 9 минут 28 секунд; предусмотренные паузы суммарно занимают {pauses} секунды. Полученное плановое время - 10 минут. Фактическое проговаривание Михаилом Антоновым не измерялось; индивидуальная репетиция остается необходимой для подтверждения регламента.')
p(report,'Первые 15,4 секунды по модели посвящены исходному противоречию. В панели докладчика таймер показывает фактическое время, плановое окно текущего шага и отклонение от него. Отклонение равно нулю внутри окна. Резервные темы приостанавливают таймер и не входят в расчет.')
p(report,'Проведена отдельная редактура: отсутствуют буква «ё», предложения, начинающиеся со слова «Это», длинные тире и универсальные риторические связки. Проверены переходы от принципа к общему предположению и специальному правилу, различие нормативного основания введения и фактов применения, пределы презюмирования и вспомогательная роль конструкции.')
report.add_heading('Источниковая опора',1)
p(report,'Основной текст - antonov_article_ready.doc. Файл «Антонов М.А. Статья.docx» сопоставлен с ним: более поздняя редакция того же исследования содержит преимущественно уточнения пунктуации, заключительного разграничения уровней и библиографии. Содержание сцен и речи следует основному тексту с учетом этих уточнений.')
p(report,'Авторский режим восстановлен по «Библиотеке научного письма Михаила Антонова», «Алгоритму авторского режима Михаила Антонова» и доступному авторскому контексту. Прежний PPTX использован только для проверки терминологии и ранее намеченного состава темы. Прежний дизайн и прежний текст речи не использованы как основа новой версии.')
p(report,'Доступный отчет Думэйт от 20 сентября относится к предыдущему материалу о реализации принципа добросовестности. Проверенный отчет именно для финальной статьи и нового устного текста не установлен. Его показатели не перенесены на данный комплект; оценка детектора не имитировалась.')
report.add_heading('Нормативная проверка',1)
p(report,'Проверены относимые положения части первой ГК РФ в доступной консолидированной редакции от 10 июня 2026 года и постановления Пленума Верховного Суда РФ № 7 в редакции от 22 июня 2021 года. Проверка проведена на 7 октября 2026 года. Новые судебные примеры не добавлялись.')
for law in D['laws']:
    link(report,law['name'],law['url']);p(report,law['result'],small=True)
p(report,'Подпункты 1 и 2 пункта 2 статьи 434.1 ГК РФ образуют альтернативные основания. Признаки внутри соответствующего основания должны быть установлены; одного немотивированного прекращения переговоров недостаточно. Статьи 12 и 15 и пункт 20 Пленума использованы в пределах вспомогательной функции презумпции и самостоятельности иных условий защиты.')
p(report,'Пункт 3 статьи 179 ГК РФ прямо не закрепляет предложенную автором презумпцию использования положения потерпевшего. Все шаги, посвященные этой модели, маркированы «АВТОРСКОЕ ПРЕДЛОЖЕНИЕ · DE LEGE FERENDA». Доказанный комплекс оснований и предполагаемый элемент показаны разными графическими обозначениями.')
report.add_heading('Технические результаты',1)
checks=[
('Chrome','Реально проверен Chrome 154.0.0.0 в облачном браузере. Отдельный экземпляр Microsoft Edge недоступен; совместимость с Edge не выдается за проведенное испытание.'),
('Размеры','Проверены внутренние размеры iframe 1366×768, 1920×1080, 2560×1440, 3840×2160 и 1920×1200. Всего 180 состояний, обрезаний заголовков, выхода текста и прокрутки аудитории не обнаружено. Размеры ОС и физические проекторы не эмулировались.'),
('Управление','ArrowRight, ArrowLeft, Space, Enter, PageDown и PageUp прошли проверку. Двадцать быстрых переходов завершились в ожидаемом состоянии. Физический кликер не подключался; проверены соответствующие клавиатурные события.'),
('Синхронизация','Проверены переходы из панели докладчика в аудиторию и из аудитории в панель. BroadcastChannel дополнен localStorage и postMessage для связанных окон. Один компьютер, один origin либо окна, открытые друг из друга.'),
('Таймер','Старт, увеличение до 00:21, пауза при входе в резерв, сброс в 00:00 и плановое окно проверены. Обычная перезагрузка сохраняет шаг. Параметр ?reset намеренно начинает заново и для сценического запуска не нужен.'),
('Полный экран','Облачная среда отклонила Fullscreen API; показано предусмотренное сообщение с F11. Фактический переход браузера в полный экран на ноутбуке не подтвержден в этой среде.'),
('Резерв','Проверены все десять тем, выбор цифрами 1-9 и 0, возврат Esc к основной линии. Созданы отдельные резервные страницы PDF и статического HTML.'),
('Автономность','Автономный HTML выполнен как srcdoc при внутреннем размере 1920×1080. Шрифт загружен, внешних ресурсных запросов нет, прокрутки нет. Работа при физически отключенном сетевом адаптере и поведение file:// на конкретном ноутбуке отдельно не измерялись.'),
('Движение и графика','Режим M отключает движение; системное prefers-reduced-motion поддержано CSS. G убирает перспективу. Проект не создает WebGL/WebGPU-контекстов, поэтому отключение этих API не нарушает основную композицию.'),
('Локальные ресурсы','Golos Text с кириллицей размещен локально и встроен в standalone.html; лицензия SIL OFL 1.1 включена. Нет внешних библиотек, видео, тяжелых текстур или фонового цикла отрисовки.'),
]
for name,result in checks:p(report,name,bold=True);p(report,result)
perf_file=Path('/workspace/scratch/keynote-performance-transitions.json')
if perf_file.is_file():
    perf=json.loads(perf_file.read_text()); f=perf.get('frameTimes',{});s=perf.get('stress',{})
    p(report,'Частота кадров',bold=True)
    p(report,f'В инструментальной серии {f.get("samples")} кадров: медиана {f.get("median",0):.2f} мс, 95-й процентиль {f.get("p95",0):.2f} мс, среднее {f.get("mean",0):.2f} мс, максимум {f.get("max",0):.2f} мс. Интервалы включают паузы и ограничения облачной сессии; стабильные 60 FPS на целевом ноутбуке не подтверждены.')
    p(report,f'Повторная навигация: {s.get("navigations")} переходов, число DOM-узлов после полного цикла от {s.get("domNodeMin")} до {s.get("domNodeMax")}. Постоянство DOM ограничивает риск накопления элементов, но не заменяет длительного профилирования памяти. Общий показатель layout-shift: {perf.get("layoutShiftTotal")}.')
else:
    p(report,'Частота кадров',bold=True)
    p(report,'Первичная серия из 120 кадров: средний интервал 33.46 мс, 95-й процентиль 16.80 мс. Серия содержит длительную облачную задержку. Стабильные 60 FPS на целевом ноутбуке не подтверждены. Размер сцены постоянен, обработчики регистрируются один раз, peer-набор ограничен, DOM заменяется при смене шага. Длительное профилирование heap не выполнено.')
stress_path=Path('/workspace/scratch/keynote-stress.json')
if stress_path.is_file():
    st=json.loads(stress_path.read_text())
    p(report,f'Дополнительная проверка устойчивости: {st["navigations"]} переходов за {st["durationMs"]:.1f} мс. После каждого из 40 полных циклов число элементов DOM оставалось равным {st["cycleNodeMin"]}; между состояниями диапазон {st["stateNodeMin"]}-{st["stateNodeMax"]}. UsedJSHeapSize вырос с {st["heapBefore"]} до {st["heapAfter"]} байт; без принудительного GC рост не доказывает утечку и не позволяет утверждать ее отсутствие. Длительная проверка heap на целевом устройстве не выполнялась.')
report.add_heading('Визуальная система',1)
p(report,'Композиция построена вокруг изменения правового режима: разрез типографики отмечает специальное исключение, правовые плоскости разводят уровни, перенос метки показывает доказательственный риск, внутренняя вставка занимает место вспомогательного средства, а штриховка отделяет предполагаемый элемент от доказанных оснований. Сцены сохраняют структуру при отключенном движении.')
p(report,'Палитра: off-black, теплая белая бумага, сталь и ограниченный медный сигнал авторского предложения. Основной шрифт Golos Text. В интерфейсе аудитории нет фотографий, обычных карточек или прокручиваемой страницы. Ключевая дефиниция показана отдельно как допустимое исключение из краткого экранного текста.')
p(report,'Исследовательские ориентиры: Vercel Ship, Stripe Sessions 2026, Linear 2026, редакционные работы Pentagram, подборки Awwwards и CSS Design Awards 2025, доступные материалы Apple Event и Bloomberg Graphics. Прямой просмотр NYT оказался ограничен; FT использован в пределах доступного материала, без заявления о просмотре закрытых интерактивов.')
for label,url in [('Vercel Ship','https://vercel.com/blog/designing-and-building-the-vercel-ship-conference-platform'),('Linear 2026','https://linear.app/now/behind-the-latest-design-refresh'),('Stripe Sessions 2026','https://stripe.com/nl-be/sessions/2026'),('Pentagram editorial','https://www.pentagram.com/editorial-design')]:link(report,label,url)
report.add_heading('Публикация и резервный запуск',1)
link(report,'Публичная презентация',SITE)
link(report,'Исходники и production build в отдельной ветке GitHub',REPO)
p(report,'В GitHub используется отдельная папка ulyanovsk-keynote-2026 в новой ветке ulyanovsk-keynote-20261007 репозитория ANTOVIUM/sgua. Основная ветка и файлы существующего проекта сохранены. Публичный хостинг размещается отдельно от существующего проекта.')
p(report,'В ZIP находятся исходники, dist с интерактивной и автономной версиями, папка static с 36 страницами, PDF и документы. Для простого автономного запуска: dist/standalone.html. Для более предсказуемой синхронизации браузерных окон: python dist/start-local.py. Для резерва без JavaScript: dist/static/index.html. PDF содержит каждый основной шаг и все десять ответов.')
publication_path=out/'publication.json'
if publication_path.is_file():
    publication=json.loads(publication_path.read_text())
    p(report,'Публикация подтверждена сервисом размещения: '+publication['status']+'. Публичный доступ подтвержден настройкой access_mode public. Публичная версия соответствует отправленному исходному состоянию. SHA и сведения о версии находятся в publication.json внутри ZIP.')
else:
    p(report,'Сведения о финальном статусе публикации и SHA находятся в publication.json внутри комплекта. Публичная ссылка считается опубликованной после успешного терминального статуса сервиса размещения; отчет не подменяет такую проверку.')
report.save(out/'Antonov_Testing_and_Sources.docx')

(work/'qa/timing.json').write_text(json.dumps({'written':written,'spoken':word_total,'pauses':pauses,'rate':word_total/((600-pauses)/60),'steps':steps},ensure_ascii=False,indent=2))
print(json.dumps({'docx':3,'pdf_pages':len(captures),'static_pages':len(captures),'written_words':written,'spoken_words':word_total,'pauses':pauses,'duration':cursor},ensure_ascii=False))
